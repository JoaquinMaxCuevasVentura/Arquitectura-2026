"""Arma las dos versiones del artículo (con autoría y anónima) sobre la plantilla de Astrágalo.

Lee el borrador en Markdown y los archivos de esta carpeta (resúmenes, declaración de IA,
referencias, pies de figuras y datos de autor), y escribe dos .docx con el formato de las
normas: Calibri 10 justificado con 6 pt después, resúmenes en 9, pies y referencias en 8.

    python3 armar_docx.py [BORRADOR] [PLANTILLA] [LAMINAS] [SALIDA]

Sin argumentos usa ../articulo/Borrador 0.md, la plantilla y las láminas de la carpeta del
proyecto, y escribe en ./docx/. No cambia el borrador.
"""

import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
BORRADOR = AQUI.parent / "articulo" / "Borrador 0.md"
PLANTILLA = Path("/mnt/project-files/convocatoria/Plantilla_normas_Astragalo_2024_Chicago17.docx")
LAMINAS = Path("/mnt/project-files/propuesta/laminas")
SALIDA = AQUI / "docx"

# Número de figura -> lámina de trabajo que la marca en el .docx (ver laminas/README.md).
FIGURAS = {1: "puentes.jpg", 2: "dos-jardines.jpg", 3: "vistas-comedor.jpg", 4: "escena-exotica.jpg"}
GRIS = RGBColor(0x80, 0x80, 0x80)
# En la versión anónima, los pies no dicen quién dibujó.
ANONIMO_PIES = [("Montaje del autor", "Montaje propio"), ("Dibujo del autor", "Dibujo propio"),
                ("números del autor", "números propios")]


def entre_marcas(ruta, marca):
    texto = Path(ruta).read_text(encoding="utf-8")
    return texto.split(f"<!-- {marca} -->", 1)[1].split("<!-- fin -->", 1)[0].strip()


def comillas(texto):
    """« » -> “ ”, y ‘ ’ para las comillas dentro de otras, como piden las normas (§7).

    El borrador anida al estilo español, «…“…”…»: las “ ” que están dentro de « » también bajan
    de nivel. Fuera de « », las “ ” se dejan como están.
    """
    salida, nivel = [], 0
    for c in texto:
        if c == "«" or (c == "“" and nivel > 0):
            salida.append("“" if nivel == 0 else "‘")
            nivel += 1
        elif c == "»" or (c == "”" and nivel > 0):
            nivel = max(nivel - 1, 0)
            salida.append("”" if nivel == 0 else "’")
        else:
            salida.append(c)
    return "".join(salida)


def limpiar(texto):
    texto = comillas(texto)
    texto = re.sub(r"\(Fig\. (\d+)\)", r"(Fig. \1.)", texto)
    texto = texto.replace(" ", " ")
    return texto


# Trozos en línea: **negrita** (pasa a cursiva: las normas no admiten negritas), *cursiva*,
# ~~tachado~~, `código`, [notas de trabajo] en gris y enlaces [texto](url).
EN_LINEA = re.compile(r"(\*\*.+?\*\*|~~.+?~~|\*[^*\s][^*]*?\*|`[^`]+`|\[[^\]]*\]\([^)]*\))")
NOTA = re.compile(r"\[(?:FALTA|POSIBLE|ESQUELETO)")


def notas_de_trabajo(texto):
    """Separa las notas [FALTA …], [POSIBLE …] y [ESQUELETO …], contando corchetes, porque
    una nota puede llevar otros adentro («[FALTA: … (Maturana y Varela [1973] 2004) …]»)."""
    trozos, i = [], 0
    for m in NOTA.finditer(texto):
        if m.start() < i:
            continue
        nivel, j = 0, m.start()
        while j < len(texto):
            nivel += {"[": 1, "]": -1}.get(texto[j], 0)
            j += 1
            if nivel == 0:
                break
        trozos += [(texto[i:m.start()], False), (texto[m.start():j], True)]
        i = j
    trozos.append((texto[i:], False))
    return [(s, nota) for s, nota in trozos if s]


def runs(parrafo, texto, tam=10, gris=False, cursiva=False, negrita=False):
    for parte, nota in notas_de_trabajo(limpiar(texto)):
        for trozo in EN_LINEA.split(parte):
            if trozo:
                run(parrafo, trozo, tam, gris or nota, cursiva, negrita)
    return parrafo


def run(parrafo, trozo, tam, gris, cursiva, negrita):
    estilo = dict(italic=cursiva, strike=False)
    if trozo.startswith("**"):
        trozo, estilo["italic"] = trozo[2:-2], True
    elif trozo.startswith("~~"):
        trozo, estilo["strike"] = trozo[2:-2], True
    elif trozo.startswith("`"):
        trozo = trozo[1:-1]
    elif re.match(r"\[[^\]]*\]\(", trozo):
        trozo = re.match(r"\[([^\]]*)\]", trozo).group(1)
    elif len(trozo) > 2 and trozo[0] == "*" and trozo[-1] == "*":
        trozo, estilo["italic"] = trozo[1:-1], not cursiva
    r = parrafo.add_run(trozo)
    r.font.size = Pt(tam)
    r.italic = estilo["italic"] or None
    r.bold = negrita or None
    r.font.strike = estilo["strike"] or None
    if gris:
        r.font.color.rgb = GRIS


def runs_con_notas(parrafo, texto, tam=10, **kw):
    """Como runs(), pero todo lo que va entre corchetes es nota para completar y va en gris."""
    for trozo in re.split(r"(\[[^\]]+\])", texto):
        if trozo:
            runs(parrafo, trozo, tam, gris=trozo.startswith("[") and trozo != "[…]", **kw)
    return parrafo


class Articulo:
    def __init__(self, plantilla):
        self.doc = Document(plantilla)
        cuerpo = self.doc.element.body
        for hijo in list(cuerpo):
            if hijo.tag != qn("w:sectPr"):
                cuerpo.remove(hijo)
        seccion = self.doc.sections[0]
        # El encabezado de la plantilla dice «Normas de presentación»: no va en el artículo.
        sect = seccion._sectPr
        for ref in sect.findall(qn("w:headerReference")):
            sect.remove(ref)
        for tp in sect.findall(qn("w:titlePg")):
            sect.remove(tp)
        # Márgenes de las normas (§3). La plantilla trae 1,5 cm abajo; las normas dicen 2.
        seccion.top_margin = seccion.bottom_margin = Cm(2)
        seccion.left_margin = seccion.right_margin = Cm(2.5)
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)
        rfonts = normal.element.rPr.find(qn("w:rFonts"))
        for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(atributo), "Calibri")
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def parrafo(self, texto="", tam=10, alineado=WD_ALIGN_PARAGRAPH.JUSTIFY, notas=False, **kw):
        p = self.doc.add_paragraph()
        p.alignment = alineado
        if texto:
            (runs_con_notas if notas else runs)(p, texto, tam, **kw)
        return p

    def titulo_apartado(self, texto):
        self.parrafo(texto, 10, WD_ALIGN_PARAGRAPH.LEFT, negrita=True)

    def cita_larga(self, texto):
        p = self.parrafo(texto, 10, cursiva=True)
        p.paragraph_format.left_indent = Cm(1.25)

    def figura(self, numero, pie, laminas):
        imagen = laminas / FIGURAS.get(numero, "")
        if imagen.is_file():
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(imagen), height=Cm(10))
        self.pie(pie)
        self.parrafo("[MARCADOR: lámina de trabajo; la figura final se redibuja a 16 × 10 cm como máximo, "
                     f"a 300 ppp, y se entrega también aparte como FIG{numero}.jpg.]", 8, gris=True)

    def pie(self, pie):
        etiqueta, resto = re.match(r"\*\*(.+?)\*\*\s*(.*)", pie).groups()
        p = self.parrafo("", 8)
        runs(p, etiqueta + " ", 8)
        runs(p, resto, 8)
        return p

    def tabla(self, filas):
        celdas = [[c.strip() for c in f.strip().strip("|").split("|")] for f in filas]
        celdas = [f for f in celdas if not all(re.fullmatch(r":?-+:?", c) for c in f)]
        t = self.doc.add_table(rows=len(celdas), cols=len(celdas[0]))
        t.style = None
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        bordes(t)
        for i, fila in enumerate(celdas):
            for j, texto in enumerate(fila):
                celda = t.cell(i, j)
                p = celda.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(2)
                runs(p, texto, 8, negrita=(i == 0))
        self.parrafo()

    def guardar(self, ruta, autor, titulo):
        props = self.doc.core_properties
        props.author = autor
        props.last_modified_by = autor
        props.title = titulo
        props.comments = props.subject = props.keywords = props.category = ""
        self.doc.save(ruta)
        depurar(ruta)


# Partes de la plantilla que no van en el artículo: la miniatura y las preferencias de Zotero
# (de la plantilla, con Chicago 16.ª), el encabezado y la imagen de ejemplo, que quedan sin uso.
SOBRAN = ["docProps/thumbnail.emf", "docProps/custom.xml", "word/header1.xml", "word/media/image1.jpeg"]


def depurar(ruta):
    with zipfile.ZipFile(ruta) as z:
        partes = {n: z.read(n) for n in z.namelist() if n not in SOBRAN}
    for nombre in ("_rels/.rels", "word/_rels/document.xml.rels", "[Content_Types].xml"):
        xml = partes[nombre].decode("utf-8")
        for sobra in SOBRAN:
            corto = sobra.split("/", 1)[1] if nombre.startswith("word/") else sobra
            xml = re.sub(r'<(Relationship|Override)[^>]*"/?%s"[^>]*/>' % re.escape(corto), "", xml)
        partes[nombre] = xml.encode("utf-8")
    app = partes["docProps/app.xml"].decode("utf-8")
    app = re.sub(r"<TotalTime>\d+</TotalTime>", "<TotalTime>0</TotalTime>", app)
    partes["docProps/app.xml"] = app.encode("utf-8")
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, datos in partes.items():
            z.writestr(nombre, datos)


def bordes(tabla):
    tblPr = tabla._tbl.tblPr
    b = tblPr.makeelement(qn("w:tblBorders"), {})
    for lado in ("top", "bottom", "insideH"):
        e = b.makeelement(qn(f"w:{lado}"), {qn("w:val"): "single", qn("w:sz"): "4", qn("w:color"): "808080"})
        b.append(e)
    tblPr.append(b)


def pies_de(ruta):
    pies = {}
    for linea in entre_seccion(ruta, "## Pies").splitlines():
        m = re.match(r"\*\*(Fig\.|Tabla) (\d+)\*\*", linea)
        if m:
            pies[(m.group(1), int(m.group(2)))] = linea
    return pies


def entre_seccion(ruta, titulo):
    texto = Path(ruta).read_text(encoding="utf-8")
    return texto.split(titulo + "\n", 1)[1].split("\n## ", 1)[0]


def bloques(texto):
    """Agrupa el Markdown en bloques: párrafos, tablas, citas, títulos."""
    actual, salida = [], []
    for linea in texto.splitlines():
        if not linea.strip():
            if actual:
                salida.append(actual)
            actual = []
        elif linea.startswith("#") or linea.strip() == "---":
            if actual:
                salida.append(actual)
            salida.append([linea])
            actual = []
        else:
            actual.append(linea)
    if actual:
        salida.append(actual)
    return salida


def armar(borrador, plantilla, laminas, salida, anonimo):
    md = Path(borrador).read_text(encoding="utf-8")
    titulo = re.search(r"^# (.+)$", md, re.M).group(1).strip()
    art = Articulo(plantilla)
    pies = pies_de(AQUI / "Figuras - pies, créditos y permisos.md")
    if anonimo:
        pies = {k: re.sub("|".join(a for a, _ in ANONIMO_PIES),
                          lambda m: dict(ANONIMO_PIES)[m.group(0)], v) for k, v in pies.items()}

    # Encabezado: título, autor y filiación (solo con autoría).
    art.parrafo(titulo, 12, negrita=True)
    art.parrafo()
    datos = entre_marcas(AQUI / "Datos de autor.md", "autor")
    nombre = re.search(r"^Nombre: (.+)$", datos, re.M).group(1)
    filiacion = re.search(r"^Filiación: (.+)$", datos, re.M).group(1)
    orcid = re.search(r"^ORCID: (.+)$", datos, re.M).group(1)
    if not anonimo:
        p = art.parrafo("", 10)
        runs_con_notas(p, nombre, 10, negrita=True)
        sup = p.add_run("1")
        sup.font.superscript, sup.font.size, sup.bold = True, Pt(10), True
        art.parrafo(f"1. {filiacion}. ORCID: {orcid}", 8, notas=True)
    art.parrafo()

    # Resúmenes.
    resumenes = entre_marcas(AQUI / "Resúmenes.md", "resumen")
    for parte in re.split(r"^## ", resumenes, flags=re.M)[1:]:
        lineas = [l for l in parte.strip().splitlines() if l.strip()]
        art.parrafo(lineas[0], 10, negrita=True)
        for linea in lineas[1:]:
            m = re.match(r"(Palabras clave|Key words|Palavras-chave):\s*(.*)", linea)
            if m:
                p = art.parrafo("", 10)
                runs(p, m.group(1) + ": ", 10, negrita=True)
                runs(p, m.group(2), 10)
            else:
                art.parrafo(linea, 9, notas=True)
        art.parrafo()

    # Epígrafe, que el borrador pone antes del resumen.
    epigrafe = re.search(r"^> (.+)\n> — (.+)$", md.split("## 1.", 1)[0], re.M)
    if epigrafe:
        art.cita_larga(epigrafe.group(1))
        p = art.parrafo("", 10)
        p.paragraph_format.left_indent = Cm(1.25)
        runs(p, "— " + epigrafe.group(2), 10)

    # Cuerpo, del apartado 1 a la declaración (sin incluirla: va desde su propio archivo).
    cuerpo = "## 1." + md.split("## 1.", 1)[1].split("## Declaración", 1)[0]
    for bloque in bloques(cuerpo):
        primera = bloque[0]
        if primera.strip() == "---":
            continue
        if primera.startswith("### "):
            m = re.match(r"### (\d+\.\d+)\.? (.+)", primera)
            art.titulo_apartado(f"{m.group(1)}. {m.group(2)}" if m else primera[4:])
        elif primera.startswith("## "):
            art.titulo_apartado(primera[3:].strip())
        elif primera.startswith("|"):
            art.tabla(bloque)
        elif primera.startswith(">"):
            lineas = [l.lstrip("> ").strip() for l in bloque]
            for linea in lineas:
                if linea.startswith("— "):
                    p = art.parrafo("", 10)
                    p.paragraph_format.left_indent = Cm(1.25)
                    runs(p, linea, 10)
                else:
                    art.cita_larga(linea)
        else:
            texto = " ".join(l.strip() for l in bloque)
            fig = re.match(r"\*\*Fig\. (\d+)\*\*", texto)
            tab = re.match(r"\*\*Tabla (\d+)\.?\*\*", texto)
            if fig:
                n = int(fig.group(1))
                art.figura(n, pies.get(("Fig.", n), texto), Path(laminas))
            elif tab:
                n = int(tab.group(1))
                art.pie(pies.get(("Tabla", n), texto))
            elif texto.startswith("["):
                art.parrafo(texto.replace("[", "", 1).rsplit("]", 1)[0].join("[]"), 10, gris=True)
            else:
                art.parrafo(texto)

    # Declaración de uso de IA.
    art.titulo_apartado("Declaración de uso de inteligencia artificial")
    for linea in entre_marcas(AQUI / "Declaración de uso de IA.md", "declaracion").splitlines():
        if linea.strip():
            art.parrafo(linea.strip(), notas=True)

    # Referencias.
    art.titulo_apartado("Referencias")
    for linea in md.split("## Referencias", 1)[1].splitlines():
        if linea.strip():
            art.parrafo(linea.strip(), 8)

    # Breve CV, solo con autoría (normas §10).
    if not anonimo:
        art.titulo_apartado("Breve CV")
        cv = datos.split("## CV", 1)[1]
        for linea in cv.strip().splitlines():
            if linea.strip():
                art.parrafo(linea.strip(), 8, notas=True)

    cuatro = " ".join(re.sub(r"[^\w\s]", "", titulo).split()[:4])
    apellido = nombre.split(",")[0].strip()
    archivo = f"{cuatro}_anonimo.docx" if anonimo else f"{cuatro}_{apellido}.docx"
    Path(salida).mkdir(parents=True, exist_ok=True)
    art.guardar(Path(salida) / archivo, "" if anonimo else re.sub(r"\s*\[.*?\]", "", nombre), titulo)
    return Path(salida) / archivo


if __name__ == "__main__":
    args = sys.argv[1:] + [None] * 4
    borrador = args[0] or BORRADOR
    plantilla = args[1] or PLANTILLA
    laminas = args[2] or LAMINAS
    salida = args[3] or SALIDA
    for anonimo in (False, True):
        print(armar(borrador, plantilla, laminas, salida, anonimo))
